"""task-agentのCloudWatch Logs(agent_invocation_cost)から、日別コスト・処理件数を
集計してExcelに出力する。handler.py(Lambda)から使われる他、単体でも実行できる。

シート:
  日別 : 日(JST)ごとのコスト(USD)・処理件数
  月別 : 月ごとの合計・日平均(日別シートへの数式)

使い方:
  python cost_report.py [--days 90] [--out cost_report.xlsx] [--profile NAME]

Logs Insightsのbin(1d)はUTC基準で日本の日付とずれるため、1時間単位で取得して
JSTの日付に集計し直す。
"""

import argparse
import time
from collections import defaultdict
from datetime import datetime, timedelta, timezone

import boto3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REGION = "ap-northeast-1"
LOG_GROUP = (
    "arn:aws:logs:ap-northeast-1:155830630328:log-group:"
    "/aws/bedrock-agentcore/runtimes/task_agent-lAAoax7UkB-DEFAULT"
)
JST = timezone(timedelta(hours=9))

QUERY = r"""
fields @timestamp, @message
| filter @message like /"metric": "agent_invocation_cost"/
| parse @message '"totalCostUsd": *,' as totalCostUsd
| parse @message '"processedFiles": *,' as processedFiles
| stats sum(totalCostUsd) as costUsd, sum(processedFiles) as files by bin(1h) as hourBin
| sort hourBin
| limit 10000
"""


def run_query(logs, days: int) -> list[dict]:
    end = datetime.now(timezone.utc)
    start = end - timedelta(days=days)
    query_id = logs.start_query(
        logGroupIdentifiers=[LOG_GROUP],
        startTime=int(start.timestamp()),
        endTime=int(end.timestamp()),
        queryString=QUERY,
        limit=10000,
    )["queryId"]
    while True:
        res = logs.get_query_results(queryId=query_id)
        if res["status"] in ("Complete", "Failed", "Cancelled", "Timeout"):
            break
        time.sleep(2)
    if res["status"] != "Complete":
        raise RuntimeError(f"Logs Insights query {res['status']}")
    return [{f["field"]: f["value"] for f in row} for row in res["results"]]


def to_float(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def aggregate_daily(rows: list[dict]) -> dict:
    """時間単位の行をJSTの日付に集計する。"""
    daily = defaultdict(lambda: [0.0, 0])
    for r in rows:
        ts = datetime.strptime(r["hourBin"], "%Y-%m-%d %H:%M:%S.%f").replace(tzinfo=timezone.utc)
        day = ts.astimezone(JST).date()
        daily[day][0] += to_float(r.get("costUsd"))
        daily[day][1] += int(to_float(r.get("files")))
    return daily


def write_excel(daily: dict, path: str) -> None:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")

    ws = wb.active
    ws.title = "日別"
    ws.append(["日付", "コスト(USD)", "処理件数"])
    for day in sorted(daily):
        cost, files = daily[day]
        ws.append([day, round(cost, 4), files])
    last = ws.max_row
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
    for r in range(2, last + 1):
        ws.cell(r, 1).number_format = "yyyy/mm/dd"
        ws.cell(r, 2).number_format = "$#,##0.0000"
    ws.freeze_panes = "A2"

    wm = wb.create_sheet("月別")
    wm.append(["月", "コスト(USD)", "処理件数", "稼働日数", "日平均コスト(USD)"])
    months = sorted({d.strftime("%Y/%m") for d in daily})
    for i, m in enumerate(months, start=2):
        # 日付列(A)を月ごとに突き合わせるため、日別シートに「yyyy/mm」を作らず
        # 範囲比較(月初以上・翌月初未満)のSUMIFSで集計する。
        y, mo = map(int, m.split("/"))
        first = datetime(y, mo, 1).date()
        nxt = datetime(y + (mo == 12), mo % 12 + 1, 1).date()
        crit = f'">="&DATE({first.year},{first.month},1),日別!$A$2:$A${last},"<"&DATE({nxt.year},{nxt.month},1)'
        wm.append([
            m,
            f"=SUMIFS(日別!$B$2:$B${last},日別!$A$2:$A${last},{crit})",
            f"=SUMIFS(日別!$C$2:$C${last},日別!$A$2:$A${last},{crit})",
            f"=COUNTIFS(日別!$A$2:$A${last},{crit})",
            f"=IF(D{i}=0,0,B{i}/D{i})",
        ])
        wm.cell(i, 2).number_format = "$#,##0.00"
        wm.cell(i, 5).number_format = "$#,##0.0000"
    n = len(months) + 1
    wm.append(["合計", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", ""])
    wm.cell(n + 1, 2).number_format = "$#,##0.00"
    for cell in wm[1]:
        cell.font, cell.fill = header_font, header_fill
    for cell in wm[n + 1]:
        cell.font = Font(bold=True)
    wm.freeze_panes = "A2"

    for sheet in (ws, wm):
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--days", type=int, default=90, help="集計対象の日数(デフォルト90。月別を見るなら31以上)")
    parser.add_argument("--out", default="cost_report.xlsx")
    parser.add_argument("--profile", default=None)
    args = parser.parse_args()

    session = boto3.Session(profile_name=args.profile, region_name=REGION)
    rows = run_query(session.client("logs"), args.days)
    daily = aggregate_daily(rows)
    write_excel(daily, args.out)
    print(f"{len(daily)}日分 -> {args.out}")


if __name__ == "__main__":
    main()
